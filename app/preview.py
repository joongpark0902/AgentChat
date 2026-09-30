"""산출물 미리보기: 파일 종류별로 화면에 그릴 수 있는 형태(JSON)로 바꾼다. 원본 파일은 읽기만 한다."""
from __future__ import annotations

import csv
import io
import mimetypes
from pathlib import Path

TEXT_LIMIT = 300_000
MAX_ROWS, MAX_COLS = 200, 40
TEXT_EXT = {".txt", ".py", ".js", ".ts", ".json", ".yaml", ".yml", ".toml", ".ini", ".cfg", ".log", ".html", ".css",
            ".xml", ".sql", ".bat", ".cmd", ".ps1", ".sh", ".java", ".c", ".cpp", ".h", ".cs", ".go", ".rs", ".r"}


def _decode(raw: bytes) -> str:
    for enc in ("utf-8-sig", "cp949"):
        try:
            return raw.decode(enc)
        except UnicodeDecodeError:
            continue
    return raw.decode("utf-8", errors="replace")


def _cell(v) -> str:
    if v is None:
        return ""
    if isinstance(v, float):
        return f"{v:,.2f}".rstrip("0").rstrip(".") if abs(v) >= 1000 else (f"{v:g}")
    if isinstance(v, int) and not isinstance(v, bool):
        return f"{v:,}"
    return str(v)


def preview_xlsx(p: Path) -> dict:
    from openpyxl import load_workbook

    # read_only: 읽기만 함(저장하지 않음 → 열려 있는 파일도 안전).
    # 값(data_only)과 수식을 함께 읽어, 계산 결과가 저장돼 있지 않은 수식 칸(엑셀로 한 번도 안 연 파일)은 수식을 보여준다.
    wb = load_workbook(p, read_only=True, data_only=True)
    wf = load_workbook(p, read_only=True, data_only=False)
    sheets = []
    try:
        for ws, wsf in zip(wb.worksheets, wf.worksheets):
            if getattr(ws, "sheet_state", "visible") != "visible":
                continue
            rows, truncated = [], False
            for i, (row, frow) in enumerate(zip(ws.iter_rows(values_only=True), wsf.iter_rows(values_only=True))):
                if i >= MAX_ROWS:
                    truncated = True
                    break
                cells = [_cell(v) if v is not None or not (isinstance(f, str) and f.startswith("="))
                         else str(f) for v, f in zip(row[:MAX_COLS], frow[:MAX_COLS])]
                if len(row) > MAX_COLS:
                    truncated = True
                rows.append(cells)
            while rows and not any(rows[-1]):
                rows.pop()
            width = max((len(r) - next((i for i, c in enumerate(reversed(r)) if c), len(r)) for r in rows), default=0)
            sheets.append({"name": ws.title, "rows": [r[:width] for r in rows], "truncated": truncated})
    finally:
        wb.close()
        wf.close()
    return {"kind": "table", "sheets": sheets}


def preview_csv(p: Path) -> dict:
    text = _decode(p.read_bytes()[:TEXT_LIMIT])
    try:
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=",\t;|")
    except csv.Error:
        dialect = csv.excel
    rows, truncated = [], False
    for i, r in enumerate(csv.reader(io.StringIO(text), dialect)):
        if i >= MAX_ROWS:
            truncated = True
            break
        rows.append(r[:MAX_COLS])
    return {"kind": "table", "sheets": [{"name": p.name, "rows": rows, "truncated": truncated}]}


def preview_docx(p: Path) -> dict:
    import docx

    d = docx.Document(str(p))
    blocks = []
    body = d.element.body
    tables = {t._element: t for t in d.tables}
    paras = {para._element: para for para in d.paragraphs}
    for child in body.iterchildren():  # 본문 순서대로 문단·표
        if child in paras:
            para = paras[child]
            if para.text.strip():
                style = (para.style.name or "").lower() if para.style is not None else ""
                blocks.append({"type": "h" if style.startswith(("heading", "title", "제목")) else "p", "text": para.text})
        elif child in tables:
            t = tables[child]
            blocks.append({"type": "table", "rows": [[c.text for c in row.cells][:MAX_COLS] for row in t.rows[:MAX_ROWS]]})
    return {"kind": "doc", "blocks": blocks}


def preview_pptx(p: Path) -> dict:
    from pptx import Presentation

    prs = Presentation(str(p))
    slides = []
    for i, s in enumerate(prs.slides, 1):
        texts = []
        for sh in s.shapes:
            if sh.has_text_frame and sh.text_frame.text.strip():
                texts.append(sh.text_frame.text.strip())
            if getattr(sh, "has_table", False) and sh.has_table:
                texts.append("\n".join(" | ".join(c.text for c in r.cells) for r in sh.table.rows))
        slides.append({"no": i, "texts": texts})
    return {"kind": "slides", "slides": slides}


def preview(p: Path, raw_url: str) -> dict:
    ext = p.suffix.lower()
    mime = mimetypes.guess_type(p.name)[0] or ""
    try:
        if mime.startswith("image/"):
            return {"kind": "image", "url": raw_url}
        if ext == ".pdf":
            return {"kind": "pdf", "url": raw_url}
        if ext in (".xlsx", ".xlsm"):
            return preview_xlsx(p)
        if ext in (".csv", ".tsv"):
            return preview_csv(p)
        if ext == ".docx":
            return preview_docx(p)
        if ext == ".pptx":
            return preview_pptx(p)
    except Exception as e:  # 미리보기 실패는 원문 대신 사유를 보여줌
        return {"kind": "error", "message": f"미리보기를 만들지 못했습니다: {type(e).__name__}: {e}"}
    raw = p.read_bytes()[:TEXT_LIMIT + 1]
    if b"\x00" in raw[:4096] and ext not in TEXT_EXT:
        return {"kind": "binary", "size": p.stat().st_size}
    text = _decode(raw[:TEXT_LIMIT])
    return {"kind": "markdown" if ext in (".md", ".markdown") else "text", "content": text,
            "truncated": len(raw) > TEXT_LIMIT, "ext": ext}
