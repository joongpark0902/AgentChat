"""v8: 사용자 결정 지적은 라운드를 돌리지 않고 넘김 / 파일을 열지 않은 승인 무효 / 컨텍스트 크기 / 결과물 덤프 도구."""
import json
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path

from app import prompts, usage
from app.config import ROOT, load_config
from app.runners import ClaudeRunner, LineContext, RunResult
from app.verdict import REVIEW_SCHEMA, parse_review
from tests.test_flow import Base, FakeRunner, review


def user_issue(desc):
    return {"verdict": "REVISE", "summary": "결정 필요", "checks": "", "requirements": [],
            "issues": [{"severity": "blocker", "description": desc, "needs_user": True}]}


class NeedsUserFlow(Base):
    async def test_user_decision_stops_rounds(self):
        FakeRunner.scripts = {"jake": ["v1"], "clara": [review("APPROVE")], "quinn": [user_issue("숨긴 열을 뺄지 정해 주세요")]}
        room = self.room()
        await self.run_to_end(room, "표 작성")
        self.assertEqual(len(self.called("jake")), 1)  # 실무자에게 다시 돌리지 않음
        self.assertTrue(room.messages[-1].get("escalation"))
        self.assertIn("숨긴 열을 뺄지", room.messages[-2]["text"])
        self.assertFalse(room.running)

    async def test_mixed_blockers_still_go_to_worker(self):
        mixed = user_issue("숨긴 열")
        mixed["issues"].append({"severity": "blocker", "description": "합계 오류", "needs_user": False})
        FakeRunner.scripts = {"jake": ["v1", "v2"], "clara": [review("APPROVE"), review("APPROVE")],
                              "quinn": [mixed, user_issue("숨긴 열")]}
        room = self.room()
        await self.run_to_end(room, "표 작성")
        self.assertEqual(len(self.called("jake")), 2)
        self.assertIn("임의로 고치지 말고", self.called("jake")[1])
        self.assertTrue(room.messages[-1].get("escalation"))

    async def test_round2_prompt_is_scoped(self):
        FakeRunner.scripts = {"jake": ["v1", "v2"], "clara": [review("APPROVE"), review("APPROVE")],
                              "quinn": [review("REVISE", ("blocker", "합계 오류")), review("APPROVE")]}
        room = self.room()
        await self.run_to_end(room, "표 작성")
        self.assertNotIn("라운드 2 이상", self.called("quinn")[0])
        self.assertIn("라운드 2 이상", self.called("quinn")[1])
        self.assertIn("needs_user", self.called("quinn")[0])


class BlindRunner(FakeRunner):
    """도구를 한 번도 쓰지 않고 답하는 감독(보고문만 읽고 승인)."""
    blind = {"clara"}

    async def run(self, prompt, workspace, **kw):
        if self.agent.key == "jake":
            (Path(workspace) / "out.txt").write_text(str(len(FakeRunner.calls)), encoding="utf-8")
        res = await super().run(prompt, workspace, **kw)
        if self.agent.key != "jake":
            res.tool_calls = 0 if self.agent.key in BlindRunner.blind else 3
        return res


class BlindApproval(Base):
    async def test_blind_approval_is_retried_then_voided(self):
        FakeRunner.scripts = {"jake": ["v1"], "clara": [review("APPROVE"), review("APPROVE")], "quinn": [review("APPROVE")]}
        from app.orchestrator import Room
        room = Room(self.cfg, self.bc, runner_factory=BlindRunner)
        await self.run_to_end(room, "파일 작성")
        clara_calls = self.called("clara")
        self.assertEqual(len(clara_calls), 2)
        self.assertEqual(clara_calls[1], prompts.NO_EVIDENCE_NUDGE)
        texts = [m["text"] for m in room.messages if m["sender"] == "system"]
        self.assertTrue(any("무효" in t for t in texts))
        self.assertNotIn("clara", [m["sender"] for m in room.messages if m.get("verdict")])


class VerdictV8(unittest.TestCase):
    def test_schema_is_strict(self):  # codex(OpenAI strict): 모든 속성이 required
        def check(node):
            if node.get("type") == "object":
                self.assertEqual(sorted(node["required"]), sorted(node["properties"]))
                self.assertFalse(node["additionalProperties"])
                for v in node["properties"].values():
                    check(v)
            if node.get("type") == "array":
                check(node["items"])
        check(REVIEW_SCHEMA)

    def test_needs_user_split_and_old_format(self):
        r = parse_review({"verdict": "REVISE", "summary": "s", "checks": "", "requirements": [
            {"item": "표 작성", "met": True, "evidence": "v2.pptx 슬라이드 3"}], "issues": [
            {"severity": "blocker", "description": "a", "needs_user": True},
            {"severity": "blocker", "description": "b", "needs_user": False},
            {"severity": "minor", "description": "c", "needs_user": True}]}, "")
        self.assertEqual([i["description"] for i in r.user_blockers], ["a"])
        self.assertEqual([i["description"] for i in r.worker_blockers], ["b"])
        self.assertIn("[충족] 표 작성 — v2.pptx 슬라이드 3", r.to_text())
        self.assertIn("[필수·사용자 결정] a", r.to_text())
        old = parse_review({"verdict": "REVISE", "summary": "s", "checks": "",
                            "issues": [{"severity": "blocker", "description": "x"}]}, "")
        self.assertEqual(len(old.worker_blockers), 1)  # 예전 형식(needs_user 없음)은 실무자 몫


class ContextV8(unittest.TestCase):
    def test_claude_context_and_tool_calls(self):
        cfg = load_config()
        r = ClaudeRunner(cfg.agents["clara"], cfg.settings)
        ctx = LineContext(None, lambda *a: None, None)
        line = lambda **ev: r.on_line(json.dumps(ev), ctx)
        usage_ = {"input_tokens": 10, "cache_read_input_tokens": 50_000, "cache_creation_input_tokens": 1_000, "output_tokens": 90}
        line(type="assistant", message={"id": "m1", "usage": usage_, "content": [
            {"type": "tool_use", "name": "PowerShell", "input": {"command": "python x.py"}}]})
        line(type="assistant", parent_tool_use_id="t1", message={"id": "m2", "usage": {"input_tokens": 999_999}, "content": []})
        line(type="assistant", message={"id": "m3", "usage": usage_, "content": [{"type": "tool_use", "name": "StructuredOutput", "input": {}}]})
        self.assertEqual(r._tool_calls, 1)             # 판정 정리(StructuredOutput)는 세지 않음
        self.assertEqual(r.context["used"], 51_100)    # 하위 에이전트 메시지는 컨텍스트에 넣지 않음
        r.parse_output(json.dumps({"type": "result", "session_id": "s", "result": "ok", "is_error": False,
                                   "modelUsage": {"claude-opus-5-5": {"contextWindow": 200_000, "inputTokens": 1}}}))
        self.assertEqual(r.context["window"], 200_000)
        r2 = ClaudeRunner(cfg.agents["clara"], cfg.settings)
        r2.restore(r.state())
        self.assertEqual(r2.context["used"], 51_100)

    def test_codex_context_from_rollout(self):
        with tempfile.TemporaryDirectory() as tmp:
            p = Path(tmp) / "rollout-x.jsonl"
            rows = [{"type": "event_msg", "payload": {"type": "token_count", "info": {
                "total_token_usage": {"input_tokens": 900_000}, "model_context_window": 258_400,
                "last_token_usage": {"input_tokens": n, "cached_input_tokens": n - 5, "output_tokens": 100}}}} for n in (1_000, 172_670)]
            p.write_text("\n".join(json.dumps(r) for r in rows), encoding="utf-8")
            self.assertEqual(usage.codex_context(p), (172_770, 258_400))
            self.assertEqual(usage.codex_context(Path(tmp) / "none.jsonl"), (None, None))

    def test_run_result_defaults_unknown(self):
        self.assertIsNone(RunResult(ok=True, text="").tool_calls)  # None = 알 수 없음 → 무효 처리 안 함


class DumpTool(unittest.TestCase):
    def run_tool(self, path):
        out = subprocess.run([sys.executable, str(ROOT / "tools" / "dump_office.py"), str(path)],
                             capture_output=True, text=True, encoding="utf-8")
        self.assertEqual(out.returncode, 0, out.stderr)
        return out.stdout

    def test_xlsx_shows_formula_hidden_and_merge(self):
        from openpyxl import Workbook
        with tempfile.TemporaryDirectory() as tmp:
            wb = Workbook()
            ws = wb.active
            ws.title = "판매채널"
            ws["A1"], ws["B1"], ws["C1"] = "매출", 100, "=B1*2"
            ws.merge_cells("A3:B3")
            ws.column_dimensions["C"].hidden = True
            wb.create_sheet("숨김").sheet_state = "hidden"
            p = Path(tmp) / "t.xlsx"
            wb.save(p)
            out = self.run_tool(p)
        for want in ("판매채널", "숨김(숨김)", "숨긴 열: C", "병합: A3:B3", "B1=100", "〔=B1*2〕"):
            self.assertIn(want, out)

    def test_pptx_shows_table_and_merge(self):
        from pptx import Presentation
        from pptx.util import Inches
        with tempfile.TemporaryDirectory() as tmp:
            prs = Presentation()
            s = prs.slides.add_slide(prs.slide_layouts[5])
            s.shapes.title.text = "광고비"
            t = s.shapes.add_table(2, 2, Inches(1), Inches(2), Inches(4), Inches(1)).table
            t.cell(0, 0).text = "홈쇼핑"
            t.cell(1, 1).text = "합계"
            t.cell(0, 0).merge(t.cell(0, 1))
            p = Path(tmp) / "t.pptx"
            prs.save(p)
            out = self.run_tool(p)
        for want in ("슬라이드 1", "광고비", "홈쇼핑(병합 1x2)", "〃", "합계"):
            self.assertIn(want, out)

    def test_prompt_points_to_tool(self):
        sp = load_config().agents["clara"].system_prompt(r"C:\ws")
        self.assertIn(str(ROOT / "tools" / "dump_office.py"), sp)
        self.assertNotIn("{dump_tool}", sp)


if __name__ == "__main__":
    unittest.main()
