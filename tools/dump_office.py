"""결과물(pptx·xlsx·docx)을 글자로 풀어 보여준다. 감독이 파일을 직접 열어 확인할 때 쓴다(읽기만 함, 저장하지 않음).

사용: python dump_office.py <파일> [--sheet 이름] [--slide 번호] [--max-rows 400]
- pptx: 슬라이드별 도형 글자, 표(병합 칸 표시)
- xlsx: 시트별 값과 수식, 숨긴 시트·행·열, 병합 범위
- docx: 본문 순서대로 문단과 표
"""
from __future__ import annotations

import argparse
import sys
from pathlib import Path


def dump_pptx(p: Path, only: int | None) -> None:
    from pptx import Presentation

    prs = Presentation(str(p))
    print(f"[pptx] {p.name} — 슬라이드 {len(prs.slides)}장")
    for no, s in enumerate(prs.slides, 1):
        if only and no != only:
            continue
        print(f"\n=== 슬라이드 {no} ===")
        for sh in s.shapes:
            if getattr(sh, "has_table", False) and sh.has_table:
                t = sh.table
                print(f"[표] {sh.name} — {len(t.rows)}행 x {len(t.columns)}열")
                for r in t.rows:
                    cells = []
                    for c in r.cells:
                        if c.is_spanned:
                            cells.append("〃")  # 병합돼 가려진 칸
                        else:
                            span = f"(병합 {c.span_height}x{c.span_width})" if c.is_merge_origin else ""
                            cells.append(c.text.replace("\n", " / ") + span)
                    print(" | ".join(cells))
            elif sh.has_text_frame and sh.text_frame.text.strip():
                print(f"[글] {sh.name}: {sh.text_frame.text.strip()}")
            elif sh.shape_type == 13:
                print(f"[그림] {sh.name}")


def dump_xlsx(p: Path, only: str | None, max_rows: int) -> None:
    from openpyxl import load_workbook
    from openpyxl.utils import get_column_letter

    wv = load_workbook(p, data_only=True)   # 저장된 계산 결과
    wf = load_workbook(p, data_only=False)  # 수식
    print(f"[xlsx] {p.name} — 시트: " + ", ".join(
        f"{ws.title}{'' if ws.sheet_state == 'visible' else '(숨김)'}" for ws in wf.worksheets))
    for ws in wf.worksheets:
        if only and ws.title != only:
            continue
        wsv = wv[ws.title]
        print(f"\n=== 시트 {ws.title} ({ws.sheet_state}) 범위 {ws.dimensions} ===")
        hidden_cols = [k for k, d in ws.column_dimensions.items() if d.hidden]
        hidden_rows = [k for k, d in ws.row_dimensions.items() if d.hidden]
        if hidden_cols:
            print("숨긴 열:", ", ".join(hidden_cols))
        if hidden_rows:
            print("숨긴 행:", ", ".join(map(str, hidden_rows[:60])))
        if ws.merged_cells.ranges:
            print("병합:", ", ".join(str(r) for r in list(ws.merged_cells.ranges)[:80]))
        shown = 0
        for row in ws.iter_rows():
            out = []
            for c in row:
                if c.value is None:
                    continue
                addr = f"{get_column_letter(c.column)}{c.row}"
                if isinstance(c.value, str) and c.value.startswith("="):
                    out.append(f"{addr}={wsv[addr].value!r} 〔{c.value}〕")
                else:
                    out.append(f"{addr}={c.value!r}")
            if out:
                print("  ".join(out))
                shown += 1
                if shown >= max_rows:
                    print(f"… {max_rows}행까지만 표시(--max-rows 로 늘릴 수 있음)")
                    break


def dump_docx(p: Path) -> None:
    import docx

    d = docx.Document(str(p))
    tables = {t._element: t for t in d.tables}
    paras = {x._element: x for x in d.paragraphs}
    print(f"[docx] {p.name}")
    for child in d.element.body.iterchildren():
        if child in paras and paras[child].text.strip():
            print(paras[child].text)
        elif child in tables:
            print("[표]")
            for r in tables[child].rows:
                print(" | ".join(c.text.replace("\n", " / ") for c in r.cells))


def main() -> int:
    ap = argparse.ArgumentParser(description="pptx·xlsx·docx 를 글자로 풀어 출력(읽기 전용)")
    ap.add_argument("file")
    ap.add_argument("--sheet")
    ap.add_argument("--slide", type=int)
    ap.add_argument("--max-rows", type=int, default=400)
    a = ap.parse_args()
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    p = Path(a.file)
    if not p.is_file():
        print(f"파일이 없습니다: {p}")
        return 1
    ext = p.suffix.lower()
    if ext == ".pptx":
        dump_pptx(p, a.slide)
    elif ext in (".xlsx", ".xlsm"):
        dump_xlsx(p, a.sheet, a.max_rows)
    elif ext == ".docx":
        dump_docx(p)
    else:
        print(f"지원하지 않는 형식입니다: {ext} (pptx·xlsx·xlsm·docx)")
        return 1
    return 0


if __name__ == "__main__":
    sys.exit(main())
